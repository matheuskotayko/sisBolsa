#!/usr/bin/env python3
"""
Renderiza o pptx gerado como HTML para conferencia visual.

Le as formas reais do arquivo (posicao, tamanho, texto, fonte) e reproduz em
divs posicionados na mesma escala. Nao e o renderizador do PowerPoint, mas
pega o que importa aqui: texto estourando o slide, sobreposicao e elemento
fora da area util.
"""
import base64
import html
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

AQUI = Path(__file__).resolve().parent
EMU_POL = 914400
PX_POL = 70  # escala do preview

ALINHAMENTO = {
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    PP_ALIGN.JUSTIFY: "justify",
    PP_ALIGN.LEFT: "left",
    None: "left",
}


def px(emu):
    return emu / EMU_POL * PX_POL


def imagem_data_uri(shape):
    img = shape.image
    return f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"


def render(caminho_pptx: Path, caminho_html: Path):
    prs = Presentation(str(caminho_pptx))
    larg, alt = px(prs.slide_width), px(prs.slide_height)
    blocos = []

    for idx, slide in enumerate(prs.slides, start=1):
        partes = [f'<div class="slide" style="width:{larg}px;height:{alt}px">']
        for sh in slide.shapes:
            estilo = (f"left:{px(sh.left):.1f}px;top:{px(sh.top):.1f}px;"
                      f"width:{px(sh.width):.1f}px;height:{px(sh.height):.1f}px")
            if sh.shape_type == 13:  # PICTURE
                partes.append(f'<img class="pic" style="{estilo}" src="{imagem_data_uri(sh)}">')
            elif sh.has_table:
                linhas = []
                for i, row in enumerate(sh.table.rows):
                    celulas = []
                    for cel in row.cells:
                        r = cel.text_frame.paragraphs[0].runs
                        tam = r[0].font.size.pt if r and r[0].font.size else 14
                        neg = "font-weight:700;" if (r and r[0].font.bold) else ""
                        celulas.append(
                            f'<td style="font-size:{tam * PX_POL / 72:.1f}px;{neg}">'
                            f'{html.escape(cel.text)}</td>')
                    linhas.append(f"<tr>{''.join(celulas)}</tr>")
                larguras = "".join(f'<col style="width:{px(c.width):.1f}px">' for c in sh.table.columns)
                partes.append(f'<table class="tb" style="{estilo}"><colgroup>{larguras}</colgroup>'
                              f"{''.join(linhas)}</table>")
            elif sh.has_text_frame and sh.text_frame.text.strip():
                paras = []
                for p in sh.text_frame.paragraphs:
                    if not p.runs:
                        continue
                    f0 = p.runs[0].font
                    tam = f0.size.pt if f0.size else 18
                    al = ALINHAMENTO.get(p.alignment, "left")
                    esp = f"margin-top:{p.space_before.pt if p.space_before else 0}px;"
                    ent = p.line_spacing if isinstance(p.line_spacing, float) else 1.1
                    corridos = []
                    for r in p.runs:
                        est = ""
                        if r.font.bold:
                            est += "font-weight:700;"
                        if r.font.italic:
                            est += "font-style:italic;"
                        try:
                            if r.font.color and r.font.color.rgb:
                                est += f"color:#{r.font.color.rgb};"
                        except Exception:
                            pass
                        corridos.append(f'<span style="{est}">{html.escape(r.text)}</span>')
                    paras.append(
                        f'<p style="font-size:{tam * PX_POL / 72:.1f}px;text-align:{al};'
                        f'line-height:{ent};{esp}">{"".join(corridos)}</p>')
                partes.append(f'<div class="tx" style="{estilo}">{"".join(paras)}</div>')
            else:
                partes.append(f'<div class="sh" style="{estilo}"></div>')
        partes.append(f'<div class="marca">slide {idx}</div></div>')
        blocos.append("".join(partes))

    doc = f"""<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8"><style>
      body {{ background:#5b5b5b; margin:0; padding:22px; font-family:"Times New Roman",Times,serif; }}
      .slide {{ position:relative; background:#fff; margin:0 auto 22px; overflow:hidden;
                box-shadow:0 2px 10px rgba(0,0,0,.35); }}
      .tx {{ position:absolute; overflow:visible; }}
      .tx p {{ margin:0; }}
      .pic {{ position:absolute; object-fit:contain; }}
      .sh {{ position:absolute; border-top:1.5px solid #000; }}
      .tb {{ position:absolute; border-collapse:collapse; table-layout:fixed; background:#fff; }}
      .tb td {{ border:1px solid #9a9a9a; padding:3px 6px; vertical-align:middle;
                font-family:"Times New Roman",Times,serif; }}
      .marca {{ position:absolute; left:6px; top:4px; font:11px monospace; color:#c00; }}
    </style></head><body>{''.join(blocos)}</body></html>"""
    caminho_html.write_text(doc, encoding="utf-8")
    print(f"preview: {caminho_html.name}  ({len(prs.slides._sldIdLst)} slides, {larg:.0f}x{alt:.0f}px cada)")


if __name__ == "__main__":
    origem = AQUI / (sys.argv[1] if len(sys.argv) > 1 else "SisBolsa_Apresentacao_API_ABNT.pptx")
    render(origem, AQUI / "preview.html")
